"""
Simulated KYC/AML onboarding checks: document completeness and validity,
fuzzy watchlist screening, rule-based risk rating, exception report (Excel),
and an audit log. Self-checks against hidden ground truth at the end.
"""
import json, re
from difflib import SequenceMatcher
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

TODAY = pd.Timestamp("2026-10-01")
HIGH_RISK = {"XA", "XB"}
THRESHOLD = 0.85

df = pd.read_csv("data/onboarding_clients.csv", parse_dates=["id_doc_expiry", "address_proof_date"]).fillna("")
watch = pd.read_csv("data/watchlist_synthetic.csv")["watchlist_name"].tolist()

def tokens(s):
    return " ".join(sorted(re.sub(r"[^a-z\s-]", " ", str(s).lower()).split()))
wtok = [(w, tokens(w)) for w in watch]

def screen(name):
    best, who = 0.0, ""
    t = tokens(name)
    for w, wt in wtok:
        sc = SequenceMatcher(None, t, wt).ratio()
        if sc > best: best, who = sc, w
    return round(best, 2), who

out = []
for _, r in df.iterrows():
    issues = []
    if r["id_doc_expiry"] < TODAY: issues.append("ID document expired")
    if (TODAY - r["address_proof_date"]).days > 90: issues.append("Address proof older than 90 days")
    if r["client_type"] == "Corporate" and r["ubo_declared"] != "Yes": issues.append("Corporate: UBO not declared")
    if not r["source_of_funds"]: issues.append("Source of funds missing")
    score, who = screen(r["client_name"])
    hit = score >= THRESHOLD
    if hit: issues.append("Watchlist potential match")
    risk = 0
    risk += 30 if r["country"] in HIGH_RISK else 0
    risk += 25 if r["pep_declared"] == "Yes" else 0
    risk += 10 * len([i for i in issues if i != "Watchlist potential match"])
    risk += 60 if hit else 0
    rating = "High" if risk >= 50 else "Medium" if risk >= 25 else "Low"
    action = ("Escalate to Compliance (potential watchlist match)" if hit else
              "Return to client for documents" if issues else
              "Enhanced due diligence (high-risk profile)" if rating == "High" else "Approve for onboarding")
    out.append({**r.to_dict(), "watchlist_score": score, "closest_watchlist_entry": who if hit else "",
                "issues": " | ".join(issues), "risk_score": risk, "risk_rating": rating, "recommended_action": action})
res = pd.DataFrame(out)
res["id_doc_expiry"] = res["id_doc_expiry"].dt.date; res["address_proof_date"] = res["address_proof_date"].dt.date

res[["client_id", "issues", "risk_rating", "recommended_action"]].to_csv("reports/audit_log.csv", index=False)
exc = res[res.issues != ""]
by_issue = exc["issues"].str.split(" \\| ").explode().value_counts().rename_axis("Issue type").reset_index(name="Count")
by_action = res["recommended_action"].value_counts().rename_axis("Recommended action").reset_index(name="Clients")
with pd.ExcelWriter("reports/KYC_Onboarding_Report.xlsx", engine="openpyxl") as xw:
    pd.DataFrame({"Metric": ["Clients reviewed", "Clean (approve)", "Clients with exceptions", "Watchlist escalations", "High risk rating"],
                  "Value": [len(res), int((res.recommended_action == "Approve for onboarding").sum()), len(exc),
                            int(res.issues.str.contains("Watchlist").sum()), int((res.risk_rating == "High").sum())]}).to_excel(xw, sheet_name="Summary", index=False)
    by_issue.to_excel(xw, sheet_name="Issues by Type", index=False)
    by_action.to_excel(xw, sheet_name="Actions", index=False)
    exc.to_excel(xw, sheet_name="Exception Detail", index=False)
    res.to_excel(xw, sheet_name="All Clients", index=False)
    for ws in xw.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3B57")
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(55, max(12, max(len(str(c.value or "")) for c in col) + 2))

truth = json.load(open("data/_ground_truth_for_validation.json"))
planted = {k for k, v in truth.items()}
flagged = set(exc["client_id"])
missed, extra = planted - flagged, flagged - planted
# per-issue recall
miss_detail = []
for cid, whys in truth.items():
    got = set(res.loc[res.client_id == cid, "issues"].iloc[0].split(" | "))
    for w in whys:
        if w not in got: miss_detail.append((cid, w))
print(f"Reviewed {len(res)} | exceptions {len(exc)} | watchlist escalations {int(res.issues.str.contains('Watchlist').sum())}")
print(by_issue.to_string(index=False)); print(by_action.to_string(index=False))
print(f"Planted issues detected: {sum(len(v) for v in truth.values()) - len(miss_detail)}/{sum(len(v) for v in truth.values())} | missed: {miss_detail}")
print(f"Unplanted clients flagged (potential false positives): {len(extra)} {sorted(extra)}")
